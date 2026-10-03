from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x14, 0x24, 0x3F)
BLUE = RGBColor(0x2F, 0x6B, 0xD0)
GREY = RGBColor(0x55, 0x5F, 0x70)
FONT = "Arial"


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.9)
    s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.8)
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.08
    for name, size, color, before in (("Heading 1", 15, NAVY, 14), ("Heading 2", 11.5, BLUE, 9), ("Heading 3", 10, NAVY, 6)):
        h = d.styles[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    # footer page number
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Obayashi Corporation (1802 JT) - Equity Research | 3 October 2026 | Page ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = GREY
    _field(p, "PAGE")
    return d


def _field(p, code):
    r = p.add_run()
    r.font.size = Pt(7.5)
    r.font.color.rgb = GREY
    for t, txt in (("begin", None), (None, code), ("end", None)):
        if t:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), t)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        r._r.append(e)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def _runs(p, text, size=None, color=None, bold=None):
    """Supports **bold** inline markup."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        if i % 2 == 1:
            r.bold = True
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def h1(d, t):
    d.add_heading(t, 1)


def h2(d, t):
    d.add_heading(t, 2)


def h3(d, t):
    d.add_heading(t, 3)


def para(d, t, size=None, italic=False, color=None, after=None):
    p = d.add_paragraph()
    _runs(p, t, size, color)
    if italic:
        for r in p.runs:
            r.italic = True
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullets(d, items, size=None):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.6)
        _runs(p, it, size)


def table(d, header, rows, widths=None, size=8, first_bold=True, hl_rows=(), align_num=True):
    t = d.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, htxt in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, "14243F")
        c.paragraphs[0].text = ""
        _runs(c.paragraphs[0], htxt, size, RGBColor(255, 255, 255), True)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].paragraphs[0].text = ""
            _runs(cells[i].paragraphs[0], str(v), size, None, first_bold and i == 0)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(1)
            if ri % 2 == 1:
                shade(cells[i], "F1F3F8")
            if ri in hl_rows:
                shade(cells[i], "DCE7FA")
    if not widths:
        widths = [17.0 / len(header)] * len(header)
    if widths:
        tot = sum(widths)
        widths = [w * 17.0 / tot for w in widths]
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(d, text, fill="EEF2FA"):
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    shade(c, fill)
    c.paragraphs[0].text = ""
    _runs(c.paragraphs[0], text, 9)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def image(d, path, width_cm=16.5, caption=None):
    d.add_picture(path, width=Cm(width_cm))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = para(d, caption, 7.5, True, GREY)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
