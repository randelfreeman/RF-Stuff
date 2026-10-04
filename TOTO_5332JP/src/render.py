"""Renderers: one content model -> PDF (reportlab) and DOCX (python-docx).

Content model: list of blocks, each a tuple:
  ('h1', text) ('h2', text) ('h3', text) ('p', text) ('bullets', [text,...])
  ('table', header_list, rows_list, caption_or_None, col_weights_or_None)
  ('img', path, caption, width_frac) ('callout', text) ('pagebreak',)
Inline markup: <b>..</b>, <i>..</i> only. Use &amp; for ampersand.
"""
import re, html
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                Table, TableStyle, Image, PageBreak, KeepTogether, ListFlowable, ListItem)
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FD = '/usr/share/fonts/truetype/liberation/'
pdfmetrics.registerFont(TTFont('Lib', FD + 'LiberationSans-Regular.ttf'))
pdfmetrics.registerFont(TTFont('Lib-B', FD + 'LiberationSans-Bold.ttf'))
pdfmetrics.registerFont(TTFont('Lib-I', FD + 'LiberationSans-Italic.ttf'))
pdfmetrics.registerFont(TTFont('Lib-BI', FD + 'LiberationSans-BoldItalic.ttf'))
from reportlab.pdfbase.pdfmetrics import registerFontFamily
registerFontFamily('Lib', normal='Lib', bold='Lib-B', italic='Lib-I', boldItalic='Lib-BI')

NAVY = colors.HexColor('#16243D')
BLUE = colors.HexColor('#2F6FD0')
GREY = colors.HexColor('#5E6878')
LIGHT = colors.HexColor('#F1F3F7')
RULE = colors.HexColor('#C9CED8')

S = {
    'h1': ParagraphStyle('h1', fontName='Lib-B', fontSize=16, leading=20, textColor=NAVY, spaceBefore=6, spaceAfter=8),
    'h2': ParagraphStyle('h2', fontName='Lib-B', fontSize=11.5, leading=15, textColor=BLUE, spaceBefore=8, spaceAfter=4),
    'h3': ParagraphStyle('h3', fontName='Lib-B', fontSize=9.5, leading=12, textColor=NAVY, spaceBefore=5, spaceAfter=2),
    'p': ParagraphStyle('p', fontName='Lib', fontSize=8.8, leading=12, textColor=colors.HexColor('#1C2433'), spaceAfter=4),
    'b': ParagraphStyle('b', fontName='Lib', fontSize=8.8, leading=11.6, textColor=colors.HexColor('#1C2433')),
    'cap': ParagraphStyle('cap', fontName='Lib-I', fontSize=7, leading=9, textColor=GREY, spaceBefore=2, spaceAfter=6),
    'th': ParagraphStyle('th', fontName='Lib-B', fontSize=7.2, leading=8.8, textColor=colors.white),
    'td': ParagraphStyle('td', fontName='Lib', fontSize=7.2, leading=8.8, textColor=colors.HexColor('#1C2433')),
    'call': ParagraphStyle('call', fontName='Lib', fontSize=9, leading=12.5, textColor=colors.white),
    'toc1': ParagraphStyle('toc1', fontName='Lib', fontSize=9.5, leading=14, leftIndent=0),
}


class Doc(BaseDocTemplate):
    def __init__(self, fn, title, sub, **kw):
        super().__init__(fn, pagesize=A4, leftMargin=16*mm, rightMargin=16*mm, topMargin=18*mm, bottomMargin=16*mm, title=title, author='Equity Research')
        self.rtitle, self.rsub = title, sub
        fr = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id='f')
        self.addPageTemplates([PageTemplate('cover', [fr], onPage=self.cover), PageTemplate('main', [fr], onPage=self.deco)])

    def cover(self, c, d):
        pass

    def deco(self, c, d):
        c.saveState()
        w, h = A4
        c.setStrokeColor(NAVY); c.setLineWidth(0.8)
        c.line(16*mm, h-12*mm, w-16*mm, h-12*mm)
        c.setFont('Lib-B', 7); c.setFillColor(BLUE)
        c.drawString(16*mm, h-10*mm, self.rtitle.upper())
        c.setFont('Lib', 7); c.setFillColor(GREY)
        c.drawRightString(w-16*mm, h-10*mm, self.rsub)
        c.drawString(16*mm, 9*mm, 'For professional investors only. Not investment advice. See data notes and disclaimer.')
        c.drawRightString(w-16*mm, 9*mm, 'Page %d' % d.page)
        c.restoreState()

    def afterFlowable(self, f):
        if isinstance(f, Paragraph) and f.style.name == 'h1':
            self.notify('TOCEntry', (0, f.getPlainText(), self.page))


def _pdf_text(t):
    return re.sub(r'&(?!amp;|lt;|gt;|#)', '&amp;', str(t))


def table_flow(header, rows, weights, width):
    n = len(header)
    if weights is None:
        lens = [max([len(re.sub('<[^>]+>', '', str(r[i]))) for r in rows] + [len(header[i])]) for i in range(n)]
        lens = [min(max(l, 4), 60) for l in lens]
        weights = lens
    tot = float(sum(weights))
    cw = [width * w / tot for w in weights]
    data = [[Paragraph(_pdf_text(h), S['th']) for h in header]]
    for r in rows:
        data.append([Paragraph(_pdf_text(x), S['td']) for x in r])
    t = Table(data, colWidths=cw, repeatRows=1)
    st = [('BACKGROUND', (0, 0), (-1, 0), NAVY), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
          ('GRID', (0, 0), (-1, -1), 0.3, RULE), ('TOPPADDING', (0, 0), (-1, -1), 2.2),
          ('BOTTOMPADDING', (0, 0), (-1, -1), 2.2), ('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3)]
    for i in range(1, len(data)):
        if i % 2 == 0:
            st.append(('BACKGROUND', (0, i), (-1, i), LIGHT))
    t.setStyle(TableStyle(st))
    return t


def build_pdf(fn, title, sub, cover_blocks, blocks):
    doc = Doc(fn, title, sub)
    W = doc.width
    story = []
    story.extend(cover_blocks(S, W))
    from reportlab.platypus import NextPageTemplate
    story.append(NextPageTemplate('main'))
    story.append(PageBreak())
    toc = TableOfContents(); toc.levelStyles = [S['toc1']]
    story.append(Paragraph('Contents', S['h1'].clone('tochead')))
    story.append(toc)
    story.append(PageBreak())
    for b in blocks:
        k = b[0]
        if k in ('h1', 'h2', 'h3', 'p'):
            story.append(Paragraph(_pdf_text(b[1]), S[k]))
        elif k == 'bullets':
            items = [ListItem(Paragraph(_pdf_text(x), S['b']), leftIndent=10, value='square') for x in b[1]]
            story.append(ListFlowable(items, bulletType='bullet', start='square', bulletFontSize=5, bulletColor=BLUE, leftIndent=10, bulletOffsetY=-2))
            story.append(Spacer(1, 4))
        elif k == 'table':
            _, hdr, rows, cap = b[:4]
            wts = b[4] if len(b) > 4 else None
            story.append(table_flow(hdr, rows, wts, W))
            if cap:
                story.append(Paragraph(_pdf_text(cap), S['cap']))
            else:
                story.append(Spacer(1, 6))
        elif k == 'img':
            _, path, cap, frac = b
            from reportlab.lib.utils import ImageReader
            ir = ImageReader(path); iw, ih = ir.getSize()
            w = W * frac; h = w * ih / iw
            flow = [Image(path, width=w, height=h)]
            if cap:
                flow.append(Paragraph(_pdf_text(cap), S['cap']))
            story.append(KeepTogether(flow))
        elif k == 'callout':
            t = Table([[Paragraph(_pdf_text(b[1]), S['call'])]], colWidths=[W])
            t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#2B3A55')),
                                   ('LINEBEFORE', (0, 0), (0, -1), 3, BLUE),
                                   ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
                                   ('LEFTPADDING', (0, 0), (-1, -1), 10), ('RIGHTPADDING', (0, 0), (-1, -1), 10)]))
            story.append(Spacer(1, 3)); story.append(t); story.append(Spacer(1, 7))
        elif k == 'pagebreak':
            story.append(PageBreak())
    doc.multiBuild(story)


# ---------------- DOCX ----------------
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _runs(par, text, size=None, color=None, bold=None):
    text = text.replace('<br/>', '\n')
    parts = re.split(r'(</?b>|</?i>)', text)
    b = i = False
    for p in parts:
        if p == '<b>': b = True; continue
        if p == '</b>': b = False; continue
        if p == '<i>': i = True; continue
        if p == '</i>': i = False; continue
        if not p: continue
        p = html.unescape(re.sub(r'<[^>]+>', '', p))
        r = par.add_run(p)
        r.bold = b or bool(bold); r.italic = i
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


def _shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcol)
    tcPr.append(shd)


def build_docx(fn, title, sub, cover_lines, blocks):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    for hn, sz, col in (('Heading 1', 16, RGBColor(0x16, 0x24, 0x3D)), ('Heading 2', 12, RGBColor(0x2F, 0x6F, 0xD0)), ('Heading 3', 10, RGBColor(0x16, 0x24, 0x3D))):
        hs = d.styles[hn]; hs.font.name = 'Arial'; hs.font.size = Pt(sz); hs.font.color.rgb = col; hs.font.bold = True
        hs.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    hdr = sec.header.paragraphs[0]; _runs(hdr, title + '  |  ' + sub, size=7, color=RGBColor(0x5E, 0x68, 0x78))
    # cover
    for i, (txt, sz, bold, colr) in enumerate(cover_lines):
        p = d.add_paragraph(); _runs(p, txt, size=sz, color=colr, bold=bold)
    d.add_page_break()
    usable = 21 - 3.6
    for b in blocks:
        k = b[0]
        if k == 'h1':
            d.add_heading(re.sub('<[^>]+>', '', html.unescape(b[1])), level=1)
        elif k == 'h2':
            d.add_heading(re.sub('<[^>]+>', '', html.unescape(b[1])), level=2)
        elif k == 'h3':
            d.add_heading(re.sub('<[^>]+>', '', html.unescape(b[1])), level=3)
        elif k == 'p':
            p = d.add_paragraph(); _runs(p, b[1]); p.paragraph_format.space_after = Pt(4)
        elif k == 'bullets':
            for x in b[1]:
                p = d.add_paragraph(style='List Bullet'); _runs(p, x); p.paragraph_format.space_after = Pt(1)
        elif k == 'table':
            _, hd, rows, cap = b[:4]
            wts = b[4] if len(b) > 4 else None
            n = len(hd)
            if wts is None:
                lens = [max([len(re.sub('<[^>]+>', '', str(r[i]))) for r in rows] + [len(hd[i])]) for i in range(n)]
                wts = [min(max(l, 4), 60) for l in lens]
            tot = float(sum(wts))
            t = d.add_table(rows=1 + len(rows), cols=n)
            t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for ci in range(n):
                w = Cm(usable * wts[ci] / tot)
                for ri in range(1 + len(rows)):
                    t.cell(ri, ci).width = w
            for ci, h in enumerate(hd):
                c = t.cell(0, ci); c.text = ''
                _runs(c.paragraphs[0], str(h), size=7.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True); _shade(c, '16243D')
            for ri, r in enumerate(rows):
                for ci, x in enumerate(r):
                    c = t.cell(ri + 1, ci); c.text = ''
                    _runs(c.paragraphs[0], str(x), size=7.5)
                    if ri % 2 == 1: _shade(c, 'F1F3F7')
            if cap:
                p = d.add_paragraph(); _runs(p, '<i>' + cap + '</i>', size=7, color=RGBColor(0x5E, 0x68, 0x78))
            else:
                d.add_paragraph()
        elif k == 'img':
            _, path, cap, frac = b
            d.add_picture(path, width=Cm(usable * frac))
            if cap:
                p = d.add_paragraph(); _runs(p, '<i>' + cap + '</i>', size=7, color=RGBColor(0x5E, 0x68, 0x78))
        elif k == 'callout':
            t = d.add_table(rows=1, cols=1); c = t.cell(0, 0); c.text = ''
            _runs(c.paragraphs[0], b[1], size=9, color=RGBColor(0xFF, 0xFF, 0xFF)); _shade(c, '2B3A55')
            d.add_paragraph()
        elif k == 'pagebreak':
            d.add_page_break()
    d.save(fn)
